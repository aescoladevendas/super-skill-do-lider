# -*- coding: utf-8 -*-
"""Diagnóstico de uma rede do atendimento direto, só com o painel de varejo (Excel padrão).

O que faz: a partir dos dois cubos (cubo_rede.json e cubo_lojas.json) monta o Excel padrão
do diagnóstico de rede: rede x mercado da rede, estados, marcas, categorias, SKUs, abertura
loja a loja (onde cresce e onde cai, sortimento) e plano de ação com dono. Tudo é dirigido
pelo dado: a leitura C-level, os ganhos e as perdas de share por categoria, as linhas do
plano de ação e a aba FONTE mudam sozinhos de uma rede para outra.

Abas: LEITURA · MÊS A MÊS · MARCAS E CATEGORIAS · SKUs DA EMPRESA · LOJAS · LOJAS x SEMANA ·
LOJAS x CATEGORIA · LOJAS x SKU · CONCORRÊNCIA (só se a rede libera as marcas concorrentes)
· PLANO DE AÇÃO · FONTE.

Entradas (na --pasta):
  cubo_rede.json    gerado por painel_varejo_cubo.py
  cubo_lojas.json   gerado por painel_varejo_lojas.py
  uf_map.json       opcional: {"NOME DA LOJA": "UF"} quando nome e endereço não trazem o estado
Da config: produtos.marcas_foco (a empresa), produtos.categoria_sem_marca_da_empresa,
organizacao.gerente, saidas.identidade_visual e a seção diagnostico_rede (marcas fora da
empresa, categorias fora da unidade, ações por categoria, limiares, rótulos).
Saída: DIAGNOSTICO_<REDE>_PAINEL_VAREJO_LOJAS.xlsx (ou --saida).

Como rodar:
  python3 montar_diagnostico.py --rede "Rede Exemplo Norte" --executivo "Executivo B" \\
      --pasta ./dados_rede --lojas-desc "PA 14, TO 4"
  opcionais: --mes-fechado MM/AA (padrão: penúltimo mês do export, o último é parcial),
  --incluir-ultimo-mes, --semanas-periodo 12, --incluir-ultima-semana, --uf-padrao PA,
  --saida arquivo.xlsx, --config empresa.yaml
Depois: recalcular no LibreOffice e exigir zero erro de fórmula (ver SKILL.md).
"""
import argparse, datetime, json, os, re, sys, unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from configuracao import carregar, obter

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
UF_SIGLAS = ('AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO').split()

# [ENGINE] limiares do método; podem ser trocados em diagnostico_rede.limiares na config
LIMIARES_PADRAO = dict(
    movimento=0.03,              # rede e estados: acima de +3% cresce, abaixo de -3% cai
    perde_share_folga=0.05,      # categoria: empresa cresce 5 pontos abaixo do mercado = perde share
    ganha_share_pp=1.0,          # categoria: ganho de share acima de 1 pp = "ganha share"
    destaque_pp=0.5,             # leitura e plano: ganho ou perda de share acima de 0,5 pp
    sem_empresa_share_max=0.02,  # categoria sem presença da empresa: share abaixo de 2%...
    sem_empresa_peso_min=0.03,   # ...e peso de pelo menos 3% do mercado da rede
    loja_movimento=0.05,         # loja cresce ou cai: P2 x P1 acima de 5% para um lado
    sku_forte=0.20,              # SKU cresce ou cai forte: 20% no acumulado do ano
    sku_perdeu_lojas=2,          # SKU perdeu lojas: mais de 2 lojas a menos que no ano anterior
    sku_relevante=0.02,          # SKU relevante da marca: 2% ou mais da venda da marca
    categoria_loja=0.10,         # LOJAS x CATEGORIA: vermelho abaixo de -10%, verde acima de +10%
    gap_lojas_alerta=3,          # LOJAS x SKU: SKU sem venda em 3 ou mais lojas fica vermelho
    gap_lojas_plano=0.08,        # plano: SKU líder sem venda em 8% das lojas (mínimo 2 lojas)
)
AGREGADOS_PADRAO = ['OUTRA MARCA', 'OUTROS PRODUTOS']


def sem_acento(texto):
    return unicodedata.normalize('NFKD', texto or '').encode('ascii', 'ignore').decode().upper().strip()


def hexa(cor):
    return cor.lstrip('#').upper()


def clarear(cor, f):
    c = hexa(cor)
    return ''.join(f'{int(int(c[i:i + 2], 16) + (255 - int(c[i:i + 2], 16)) * f):02X}' for i in (0, 2, 4))


def ml(m):
    """'08/26' → 'ago/26'."""
    a, b = m.split('/')
    return f'{MESES[int(a) - 1]}/{b}'


def ptbr(t):
    """Decimal com ponto dentro de texto → vírgula (1.5 → 1,5); milhar com ponto fica."""
    return re.sub(r'(\d)\.(\d)(?!\d)', r'\1,\2', t)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--rede', required=True, help='nome da rede no material')
    ap.add_argument('--executivo', default=None, help='dono do plano (padrão: o da config)')
    ap.add_argument('--pasta', default='.')
    ap.add_argument('--saida', default=None)
    ap.add_argument('--lojas-desc', default='', help='ex.: "PA 14, TO 4"')
    ap.add_argument('--mes-fechado', default=None, help='MM/AA do último mês fechado')
    ap.add_argument('--incluir-ultimo-mes', action='store_true')
    ap.add_argument('--semanas-periodo', type=int, default=12)
    ap.add_argument('--incluir-ultima-semana', action='store_true')
    ap.add_argument('--uf-padrao', default=None)
    ap.add_argument('--config', default=None)
    a = ap.parse_args()
    cfg = carregar(a.config)

    REDE = a.rede
    EMP = obter(cfg, 'diagnostico_rede.rotulo_empresa', None) or obter(cfg, 'meta.empresa', 'Empresa')
    GERENTE = obter(cfg, 'organizacao.gerente', 'Gerente')
    EXEC = a.executivo
    if not EXEC:  # dono = executivo do cadastro cujo nome bate com a rede
        for c in obter(cfg, 'clientes.cadastro', []) or []:
            if sem_acento(c.get('nome', '')).startswith(sem_acento(REDE)):
                EXEC = c.get('executivo')
                break
    EXEC = EXEC or 'Executivo'
    MARCAS = obter(cfg, 'produtos.marcas_foco', []) or []
    NOME_MARCA = {sem_acento(m): m for m in MARCAS}
    disp = lambda m: NOME_MARCA.get(sem_acento(m), m.title())
    FORA = obter(cfg, 'diagnostico_rede.marcas_fora_da_empresa', []) or []
    SEM_MARCA = {sem_acento(x) for x in obter(cfg, 'produtos.categoria_sem_marca_da_empresa', []) or []}
    CAT_FORA = {sem_acento(x) for x in obter(cfg, 'diagnostico_rede.categorias_fora_da_unidade', []) or []}
    ACOES = obter(cfg, 'diagnostico_rede.acoes_por_categoria', {}) or {}
    ACOES = {sem_acento(k): v for k, v in ACOES.items()}
    AGREG = {sem_acento(x) for x in (obter(cfg, 'diagnostico_rede.rotulos_concorrencia_agregada', None)
                                     or AGREGADOS_PADRAO)}
    LIM = dict(LIMIARES_PADRAO)
    LIM.update(obter(cfg, 'diagnostico_rede.limiares', {}) or {})

    C = json.load(open(os.path.join(a.pasta, 'cubo_rede.json'), encoding='utf-8'))
    LJ = json.load(open(os.path.join(a.pasta, 'cubo_lojas.json'), encoding='utf-8'))
    arq_uf = os.path.join(a.pasta, 'uf_map.json')
    UFMAP = json.load(open(arq_uf, encoding='utf-8')) if os.path.exists(arq_uf) else {}
    TRECHOS = obter(cfg, 'diagnostico_rede.uf_por_trecho', {}) or {}

    # ---------- visual (tokens da skill padrao-visual em saidas.identidade_visual)
    iv = obter(cfg, 'saidas.identidade_visual', {}) or {}
    DEST = hexa(iv.get('destaque', '#FF007F'))
    GRAF = hexa(iv.get('grafite', '#1D1D1F'))
    CINZA_T = hexa(iv.get('cinza_texto', '#6E6E73'))
    FONTE = iv.get('fonte_office', 'Arial')
    F = lambda **k: Font(name=FONTE, **k)
    sol = lambda c: PatternFill('solid', fgColor=c)
    HFILL, TFILL = sol(GRAF), sol(DEST)
    EMPF = sol(clarear(DEST, 0.88))                        # linha da rede e marcas da empresa
    GF = sol(hexa(iv.get('fundo', '#F5F5F7')))             # cinza: faixa e gap de sortimento
    VF = sol(clarear(iv.get('verde', '#248A3D'), 0.85))
    XF = sol(clarear(iv.get('vermelho', '#D70015'), 0.88))
    fino = Side(style='thin', color='D9D9D9')
    BORDA = Border(top=fino, bottom=fino, left=fino, right=fino)
    BRL, PCT, INT = '#,##0', '0.0%', '#,##0'
    PP = '+0.0" pp";-0.0" pp";0.0" pp"'
    VAR = '+0.0%;-0.0%;0.0%'

    def titulo(ws, t, sub=None, w=14):
        ws['A1'] = t
        ws['A1'].font = F(bold=True, size=16, color='FFFFFF')
        for c in range(1, w + 1):
            ws.cell(1, c).fill = TFILL
        ws.row_dimensions[1].height = 28
        if sub:
            ws['A2'] = sub
            ws['A2'].font = F(italic=True, size=9, color=CINZA_T)

    def cabecalho(ws, r, cols):
        for j, h in enumerate(cols):
            c = ws.cell(r, 1 + j, h)
            c.font = F(bold=True, color='FFFFFF', size=9)
            c.fill = HFILL
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = BORDA
        ws.row_dimensions[r].height = 30

    def cel(ws, r, c, v, fmt=None, bold=False, fill=None, wrap=False):
        x = ws.cell(r, c, v)
        x.font = F(bold=bold, size=9)
        x.border = BORDA
        if fmt:
            x.number_format = fmt
        if fill:
            x.fill = fill
        if wrap:
            x.alignment = Alignment(wrap_text=True, vertical='top')
        return x

    def faixa(ws, r, t, w):
        ws.cell(r, 1, t).font = F(bold=True, size=10, color=GRAF)
        for c in range(1, w + 1):
            ws.cell(r, c).fill = GF

    def paragrafo(ws, r, t, w=14, h=34):
        ws.cell(r, 1, ptbr(t)).font = F(size=10)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=w)
        ws.row_dimensions[r].height = h

    def larguras(ws, w):
        for i, x in enumerate(w):
            ws.column_dimensions[L(i + 1)].width = x

    S = lambda arr, idx: sum(arr[i] for i in idx)
    var = lambda x, y: (x / y - 1) if y else None
    share = lambda x, y: (x / y) if y else None
    BR = lambda v: f'{v:,.0f}'.replace(',', '.')
    sp = lambda x: f'{x * 100:+.1f}%' if x is not None else 'sem base'
    mi = lambda v: f'{v / 1e6:.1f}'

    # ---------- calendário: mês fechado, acumulado do ano e últimos 3 meses
    M = C['REDE']['months']
    NM = len(M)
    ix = {m: i for i, m in enumerate(M)}
    fech = a.mes_fechado or (M[-1] if a.incluir_ultimo_mes else M[-2])
    if fech not in ix:
        sys.exit(f'Mês fechado {fech} não está no export ({M[0]} a {M[-1]})')
    mm, aa = int(fech[:2]), int(fech[3:])
    la = aa - 1
    exigidos = [f'{m:02d}/{y:02d}' for y in (aa, la) for m in range(1, mm + 1)]
    faltam = [m for m in exigidos if m not in ix]
    if faltam:
        sys.exit('O export precisa cobrir jan do ano anterior até o mês fechado; faltam: ' + ', '.join(faltam))
    Y = [ix[f'{m:02d}/{aa:02d}'] for m in range(1, mm + 1)]
    YL = [ix[f'{m:02d}/{la:02d}'] for m in range(1, mm + 1)]
    A, AL = ix[fech], ix[f'{mm:02d}/{la:02d}']
    if AL < 2:
        sys.exit('O export precisa trazer pelo menos 2 meses antes do mesmo mês do ano anterior')
    ANT = A - 1
    T3, T3L = [A - 2, A - 1, A], [AL - 2, AL - 1, AL]
    FECHADOS = list(range(0, A + 1))
    PARCIAL = M[A + 1] if A + 1 < NM else None
    ACUM = f'jan-{MESES[mm - 1]}' if mm > 1 else 'jan'
    ANO, ANO_ANT = f'{aa:02d}', f'{la:02d}'
    ANO4, ANO4_ANT = 2000 + aa, 2000 + la
    T3TXT = f'{MESES[int(M[T3[0]][:2]) - 1]}-{MESES[mm - 1]}'

    def mes_mais(k):
        t = ANO4 * 12 + (mm - 1) + k
        return f'{MESES[t % 12]}/{(t // 12) % 100:02d}'

    def mov(dv, dpp):
        if dv is None:
            return 'sem base'
        if dv > LIM['movimento'] and (dpp or 0) > 0:
            return 'cresce e ganha share'
        if dv > LIM['movimento']:
            return 'cresce, perde share'
        if dv < -LIM['movimento'] and (dpp or 0) < 0:
            return 'cai e perde share'
        if dv < -LIM['movimento']:
            return 'cai, ganha share'
        return 'estável'

    # ---------- semanas da abertura por loja: P1 x P2 (blocos iguais de semanas fechadas)
    W = LJ['weeks']
    nW = len(W)
    fechadas = list(range(nW if a.incluir_ultima_semana else nW - 1))
    n = min(a.semanas_periodo, len(fechadas) // 2)
    if n < 1:
        sys.exit('Semanas insuficientes na abertura por loja')
    P2 = fechadas[len(fechadas) - n:]
    P1 = fechadas[len(fechadas) - 2 * n:len(fechadas) - n]
    FULL = P1 + P2
    NSEM = len(FULL)
    sem = lambda w: (re.match(r'(\d{2}/\d{2})/\d{4}$', w).group(1) if re.match(r'\d{2}/\d{2}/\d{4}$', w) else w)
    P1TXT = f'{sem(W[P1[0]])} a {sem(W[P1[-1]])}'
    P2TXT = f'{sem(W[P2[0]])} a {sem(W[P2[-1]])}'
    FULLTXT = f'{sem(W[FULL[0]])} a {W[FULL[-1]]}'
    lojas = LJ['lojas']
    NL = len(lojas)

    # ---------- unidades (rede primeiro, estados por tamanho, a derivada logo depois da origem)
    ufs = sorted([t for t in C if t.startswith('UF_') and t != 'UF_DEMAIS'],
                 key=lambda t: -S(C[t]['emp']['v'], Y))
    UNIDADES = ['REDE']
    for t in ufs:
        UNIDADES.append(t)
        if C.get('UF_DEMAIS', {}).get('derivada_de') == t:
            UNIDADES.append('UF_DEMAIS')
    exportados = [t[3:] for t in ufs]
    UF_PADRAO = a.uf_padrao or (exportados[0] if len(exportados) == 1 else 'N/D')

    LJM = C['REDE']['lojas']
    L_ant = [LJM[ix[f'{m:02d}/{la:02d}']] for m in range(1, 13) if f'{m:02d}/{la:02d}' in ix]
    L_atu = [LJM[i] for i in Y]
    L3, L3L = max(LJM[i] for i in T3), max(LJM[i] for i in T3L)
    fx = lambda l: str(min(l)) if min(l) == max(l) else f'{min(l)} a {max(l)}'
    COMPTXT = (f'Base comparável: em {ANO4_ANT} a rede tinha {fx(L_ant)} lojas com venda no painel de '
               f'varejo e em {ANO4} tem {fx(L_atu)}. A comparação {T3TXT} ({L3L} x {L3} lojas) na seção 2 '
               f"é a leitura de {'mesma base' if L3L == L3 else 'base quase comparável'}"
               + ('' if min(L_ant) == max(L_atu) else f'; o acumulado {ACUM} carrega a diferença de lojas entre os anos')
               + '.')
    HASCONC = any(sem_acento(k.split('|')[1]) not in AGREG for k in C['REDE']['conc'])

    def ltot(d, idx, cat=None):
        return sum(s['v'][i] for s in d['skus'].values() if (cat is None or s['cat'] == cat) for i in idx)

    RX_FIM = re.compile(r'(?:-|/|\()\s*([A-Z]{2})\s*\)?\s*$')

    def uf_loja(nome):
        nome = nome.strip()
        if nome in UFMAP:
            return UFMAP[nome]
        end = lojas.get(nome, {}).get('end', '')
        for texto in (sem_acento(nome), sem_acento(end)):
            m = RX_FIM.search(texto)
            if m and m.group(1) in UF_SIGLAS:
                return m.group(1)
        for trecho, uf in TRECHOS.items():
            if sem_acento(trecho) in sem_acento(nome) or sem_acento(trecho) in sem_acento(end):
                return uf
        return UF_PADRAO

    RU = sem_acento(REDE)
    _tw = lambda w: w if (w[:2] == 'LJ' and w[2:].isdigit()) or w.isdigit() or w in UF_SIGLAS else w.title()

    def SN(k):  # nome curto da loja: tira o prefixo da rede e arruma a caixa
        k = k.strip()
        resto = k[len(RU):].strip() if sem_acento(k).startswith(RU) else k
        resto = ' '.join(_tw(w) for w in resto.split())
        return f'{REDE} {resto}'.strip() if sem_acento(k).startswith(RU) else resto

    SNS = lambda k: ((SN(k)[len(REDE):].strip() if SN(k).startswith(REDE) else SN(k)) or SN(k))[:18]
    vale_cat = lambda k: sem_acento(k) not in CAT_FORA

    # =====================================================================================
    wb = Workbook()
    ws = wb.active
    ws.title = 'LEITURA'
    c = C['REDE']
    m26, m25 = S(c['mkt']['v'], Y), S(c['mkt']['v'], YL)
    r26, r25 = S(c['emp']['v'], Y), S(c['emp']['v'], YL)
    titulo(ws, f'{REDE.upper()} · DIAGNÓSTICO DO PAINEL DE VAREJO · sell out nas lojas · P{mm:02d} fechado '
               f'({ACUM}/{ANO} vs {ACUM}/{ANO_ANT})',
           f'Fonte única: painel de varejo (rede {REDE}, {LJM[A]} lojas em {ml(fech)}'
           f"{(': ' + a.lojas_desc) if a.lojas_desc else ''}). A empresa = {', '.join(MARCAS)}. "
           f'Mercado = todas as marcas nas categorias da empresa. Meses fechados até {ml(fech)}'
           f"{('; ' + ml(PARCIAL) + ' parcial excluído') if PARCIAL else ''}. Loja a loja: semanas de "
           f'{FULLTXT} (o portal abre lojas só a partir de {W[0]}).', 14)
    r = 4
    faixa(ws, r, f'1. FECHAMENTO · rede e estados · {ACUM} {ANO4} vs {ACUM} {ANO4_ANT}', 14)
    r += 1
    cabecalho(ws, r, ['Unidade', f'Lojas {ml(fech)}', f'Lojas {ml(M[AL])}', f'Mercado {ACUM} {ANO}',
                      f'Mercado {ACUM} {ANO_ANT}', 'Var. mercado', f'{EMP} {ACUM} {ANO}',
                      f'{EMP} {ACUM} {ANO_ANT}', f'Var. {EMP}', f'Share {ANO}', f'Share {ANO_ANT}',
                      'Δ share (pp)', f'{EMP} {ml(fech)} vs {ml(M[AL])}', 'Movimento'])
    r += 1
    for t in UNIDADES:
        u = C[t]
        a26, a25, b26, b25 = S(u['mkt']['v'], Y), S(u['mkt']['v'], YL), S(u['emp']['v'], Y), S(u['emp']['v'], YL)
        s26, s25 = share(b26, a26), share(b25, a25)
        dpp = ((s26 or 0) - (s25 or 0)) * 100
        vals = [u['name'], u['lojas'][A], u['lojas'][AL], a26, a25, var(a26, a25), b26, b25, var(b26, b25),
                s26, s25, dpp, var(u['emp']['v'][A], u['emp']['v'][AL]), mov(var(b26, b25), dpp)]
        fm = [None, INT, INT, BRL, BRL, VAR, BRL, BRL, VAR, PCT, PCT, PP, VAR, None]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j], bold=(t == 'REDE'), fill=EMPF if t == 'REDE' else None)
        r += 1
    paragrafo(ws, r, COMPTXT, 14, 32)
    r += 2
    faixa(ws, r, f'2. ÚLTIMOS 3 MESES FECHADOS ({T3TXT}/{ANO} vs {T3TXT}/{ANO_ANT}) · base '
                 f"{'comparável' if L3L == L3 else 'quase comparável'} ({L3L} x {L3} lojas)", 14)
    r += 1
    cabecalho(ws, r, ['Unidade', f'Mercado {T3TXT} {ANO}', f'Mercado {T3TXT} {ANO_ANT}', 'Var. mercado',
                      f'{EMP} {T3TXT} {ANO}', f'{EMP} {T3TXT} {ANO_ANT}', f'Var. {EMP}', f'Share {ANO}',
                      f'Share {ANO_ANT}', 'Δ share (pp)', f'{EMP} {ml(fech)}', f'{EMP} {ml(M[ANT])}',
                      f'{MESES[mm - 1].capitalize()} vs {MESES[int(M[ANT][:2]) - 1]}', ''])
    r += 1
    for t in UNIDADES:
        u = C[t]
        a26, a25, b26, b25 = S(u['mkt']['v'], T3), S(u['mkt']['v'], T3L), S(u['emp']['v'], T3), S(u['emp']['v'], T3L)
        vals = [u['name'], a26, a25, var(a26, a25), b26, b25, var(b26, b25), share(b26, a26), share(b25, a25),
                ((share(b26, a26) or 0) - (share(b25, a25) or 0)) * 100, u['emp']['v'][A], u['emp']['v'][ANT],
                var(u['emp']['v'][A], u['emp']['v'][ANT])]
        fm = [None, BRL, BRL, VAR, BRL, BRL, VAR, PCT, PCT, PP, BRL, BRL, VAR]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j], bold=(t == 'REDE'), fill=EMPF if t == 'REDE' else None)
        r += 1
    r += 1
    faixa(ws, r, f'3. CATEGORIAS NA REDE · mercado da rede x {EMP} · {ACUM} {ANO} vs {ANO_ANT}', 14)
    r += 1
    cabecalho(ws, r, ['Categoria', f'Mercado {ANO}', f'Mercado {ANO_ANT}', 'Var. mercado', f'{EMP} {ANO}',
                      f'{EMP} {ANO_ANT}', f'Var. {EMP}', f'Share {ANO}', f'Share {ANO_ANT}', 'Δ pp',
                      'Peso no mercado', 'Marcas da empresa', 'Leitura', ''])
    r += 1
    zero = {'v': [0] * NM}
    for k, x in sorted(c['cat_mkt'].items(), key=lambda kv: -S(kv[1]['v'], Y)):
        if not vale_cat(k):
            continue
        rb = c['cat_emp'].get(k, zero)
        a26, a25, b26, b25 = S(x['v'], Y), S(x['v'], YL), S(rb['v'], Y), S(rb['v'], YL)
        s26, s25 = share(b26, a26), share(b25, a25)
        dpp = ((s26 or 0) - (s25 or 0)) * 100
        mb = ', '.join(disp(m) for m, v in sorted([(kk.split('|')[1], S(vv['v'], Y)) for kk, vv in c['cat_marca'].items()
                                                   if kk.split('|')[0] == k], key=lambda z: -z[1]) if v > 0)
        if (b26 == 0 and b25 == 0) or sem_acento(k) in SEM_MARCA:
            txt, fill = 'Categoria sem marca da empresa na rede', GF
        elif var(a26, a25) is not None and var(b26, b25) is not None and var(b26, b25) < var(a26, a25) - LIM['perde_share_folga']:
            txt, fill = 'A empresa cresce abaixo do mercado da rede: perde share', XF
        elif dpp > LIM['ganha_share_pp']:
            txt, fill = 'A empresa ganha share na rede', VF
        else:
            txt, fill = 'Acompanha o mercado', None
        vals = [k, a26, a25, var(a26, a25), b26, b25, var(b26, b25) if b25 else None, s26, s25, dpp,
                share(a26, m26), mb or '-', txt]
        fm = [None, BRL, BRL, VAR, BRL, BRL, VAR, PCT, PCT, PP, PCT, None, None]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j], fill=fill if j == 12 else None)
        r += 1
    r += 1
    faixa(ws, r, '4. LEITURA C-LEVEL', 14)
    r += 1
    rk = sorted(lojas.items(), key=lambda kv: -ltot(kv[1], FULL))
    top5 = [f'{SN(k)} (R$ {BR(ltot(v, FULL))})' for k, v in rk[:5]]
    cresc = sorted([(k, ltot(v, P2), ltot(v, P1)) for k, v in lojas.items()],
                   key=lambda z: -(z[1] / z[2] - 1) if z[2] else 0)
    queda = [z for z in cresc if z[2] and z[1] / z[2] - 1 < -LIM['loja_movimento']]
    sobe = [z for z in cresc if z[2] and z[1] / z[2] - 1 > LIM['loja_movimento']]

    def catshare(k, idx):
        x = S(c['cat_emp'].get(k, zero)['v'], idx)
        y = S(c['cat_mkt'].get(k, zero)['v'], idx)
        return (x / y) if y else None

    CATINFO = []
    for k in c['cat_mkt']:
        if not vale_cat(k):
            continue
        s26, s25 = catshare(k, Y), catshare(k, YL)
        if s26 is None or s25 is None:
            continue
        CATINFO.append(dict(cat=k, s26=s26, s25=s25, dpp=(s26 - s25) * 100,
                            mkt26=S(c['cat_mkt'][k]['v'], Y), mkt25=S(c['cat_mkt'][k]['v'], YL),
                            rb26=S(c['cat_emp'].get(k, zero)['v'], Y), rb25=S(c['cat_emp'].get(k, zero)['v'], YL)))
    GANHOS = sorted([x for x in CATINFO if x['rb26'] > 0 and x['dpp'] > LIM['destaque_pp']], key=lambda x: -x['dpp'])
    PERDAS = sorted([x for x in CATINFO if (x['rb26'] > 0 or x['rb25'] > 0) and x['dpp'] < -LIM['destaque_pp']],
                    key=lambda x: x['dpp'])
    # categorias relevantes (peso >= 3% do mercado da rede) onde a empresa não tem presença (share < 2%)
    SEMEMP = sorted([x for x in CATINFO if x['s26'] < LIM['sem_empresa_share_max']
                     and x['mkt26'] >= LIM['sem_empresa_peso_min'] * m26], key=lambda x: -x['mkt26'])
    fg = lambda x: f"{x['cat']} ({x['s26'] * 100:.1f}%, {x['dpp']:+.1f} pp)"
    fl = lambda x: f"{x['cat']} ({x['s26'] * 100:.1f}%, era {x['s25'] * 100:.1f}%)"
    gm, gr = var(m26, m25), var(r26, r25)
    if gr is not None and gm is not None:
        ritmo = ('a empresa cresce acima da rede e ganha share' if gr > gm else
                 ('a empresa cresce, mas a rede cresce mais rápido nas nossas categorias' if gr > 0 else
                  ('a empresa cai enquanto a rede cresce' if gm > 0 else 'a empresa e a rede caem')))
    else:
        ritmo = f'sem base de comparação em {ANO4_ANT}'
    sh26, sh25 = share(r26, m26) or 0, share(r25, m25) or 0
    linhas = [
        f'Mercado das categorias da empresa dentro da rede {REDE}: R$ {mi(m26)} Mi em {ACUM}/{ANO}, {sp(gm)} vs '
        f'{ANO4_ANT}. Empresa: R$ {mi(r26)} Mi, {sp(gr)}. Share {"caiu" if sh26 < sh25 else "subiu"} de '
        f'{sh25 * 100:.1f}% para {sh26 * 100:.1f}% ({(sh26 - sh25) * 100:+.1f} pp): {ritmo}.',
        f"Onde o share sobe: {', '.join(fg(x) for x in GANHOS[:4]) if GANHOS else 'nenhuma categoria acima de +0,5 pp'}. "
        f"Onde cai: {', '.join(fl(x) for x in PERDAS[:4]) if PERDAS else 'nenhuma categoria abaixo de -0,5 pp'}."
        + (f" {SEMEMP[0]['cat']} (R$ {mi(SEMEMP[0]['mkt26'])} Mi {ACUM}/{ANO}, {sp(var(SEMEMP[0]['mkt26'], SEMEMP[0]['mkt25']))}, "
           f"{SEMEMP[0]['mkt26'] / m26 * 100:.0f}% do mercado da rede) é a maior categoria da rede sem presença da "
           f"empresa (share {SEMEMP[0]['s26'] * 100:.1f}%)." if SEMEMP else ''),
        f"Lojas ({FULLTXT}): as cinco maiores em vendas da empresa são {'; '.join(top5)}. Entre as {NL} lojas, "
        f"{len(sobe)} crescem mais de 5% no 2º período de {n} semanas contra o 1º e "
        f"{(str(len(queda)) + ' caem mais de 5%: ' + ', '.join(SN(z[0]) for z in queda[-6:])) if queda else 'nenhuma cai mais de 5%'}.",
        f'Comparável loja a loja contra {ANO4_ANT} não existe no painel de varejo (a abertura por loja começa em '
        f'{W[0]}); a leitura por loja é de tendência dentro de {ANO4} ({n} semanas x {n} semanas) e de sortimento '
        f'(SKUs com venda por loja). A comparação com {ANO4_ANT} é feita por rede e por estado.',
        f"Plano para {EXEC} (aba PLANO DE AÇÃO): 1) recuperar as categorias onde a empresa perde share "
        f"({', '.join(x['cat'] for x in PERDAS[:3]) if PERDAS else 'nenhuma em perda relevante'}); 2) atacar as lojas "
        f"em queda e as lojas com sortimento incompleto (aba LOJAS x SKU); 3) defender as categorias onde já ganhamos "
        f"share ({', '.join(x['cat'] for x in GANHOS[:3]) if GANHOS else 'manter execução'}); 4) "
        + (f"decidir com o cliente a entrada em {SEMEMP[0]['cat']}, hoje R$ {mi(SEMEMP[0]['mkt26'])} Mi com share da "
           f"empresa de {SEMEMP[0]['s26'] * 100:.1f}%." if SEMEMP else
           'revisar preço e execução com o cliente nas categorias que acompanham o mercado.')]
    for t in linhas:
        paragrafo(ws, r, t, 14, 44)
        r += 1
    larguras(ws, [30, 10, 10, 14, 14, 10, 14, 14, 10, 9, 9, 9, 12, 40])
    ws.freeze_panes = 'A4'

    # =====================================================================================
    ws = wb.create_sheet('MÊS A MÊS')
    NC = len(FECHADOS)
    titulo(ws, f'MÊS A MÊS · {ml(M[0])} a {ml(fech)} (meses fechados) · rede e estados · mercado da rede x {EMP} x share',
           'Variações vivas em fórmula. Lojas = número de lojas com venda no mês (base de comparabilidade).', NC + 3)
    r = 4
    soma_y = lambda rr: f'=SUM({L(2 + Y[0])}{rr}:{L(2 + Y[-1])}{rr})'
    soma_yl = lambda rr: f'=SUM({L(2 + YL[0])}{rr}:{L(2 + YL[-1])}{rr})'
    for t in UNIDADES:
        u = C[t]
        faixa(ws, r, u['name'], NC + 3)
        r += 1
        cabecalho(ws, r, ['Métrica'] + [ml(M[i]) for i in FECHADOS] + [f'{ACUM} {ANO}', f'{ACUM} {ANO_ANT}'])
        r += 1
        r0 = r
        for nome, arr in [('Mercado R$ (categorias da empresa)', u['mkt']['v']), (f'{EMP} R$', u['emp']['v'])]:
            cel(ws, r, 1, nome, bold=True, fill=EMPF if nome.startswith(EMP) else None)
            for i in FECHADOS:
                cel(ws, r, 2 + i, arr[i], BRL)
            cel(ws, r, NC + 2, soma_y(r), BRL, bold=True)
            cel(ws, r, NC + 3, soma_yl(r), BRL)
            r += 1
        cel(ws, r, 1, f'Share {EMP} %', bold=True)
        shr = lambda col: f'=IF({col}{r0}=0,"",{col}{r0 + 1}/{col}{r0})'
        for i in FECHADOS:
            cel(ws, r, 2 + i, shr(L(2 + i)), PCT)
        cel(ws, r, NC + 2, shr(L(NC + 2)), PCT, bold=True)
        cel(ws, r, NC + 3, shr(L(NC + 3)), PCT)
        rs = r
        r += 1
        cel(ws, r, 1, 'Lojas com venda')
        for i in FECHADOS:
            cel(ws, r, 2 + i, u['lojas'][i], INT)
        r += 1
        for nome, rr in [('Mercado vs mesmo mês do ano anterior', r0), (f'{EMP} vs mesmo mês do ano anterior', r0 + 1)]:
            cel(ws, r, 1, nome)
            for i in FECHADOS:
                cel(ws, r, 2 + i, f'=IF({L(2 + i - 12)}{rr}=0,"",{L(2 + i)}{rr}/{L(2 + i - 12)}{rr}-1)' if i >= 12 else None, VAR)
            cel(ws, r, NC + 2, f'=IF({L(NC + 3)}{rr}=0,"",{L(NC + 2)}{rr}/{L(NC + 3)}{rr}-1)', VAR, bold=True)
            r += 1
        cel(ws, r, 1, f'{EMP} vs mês anterior')
        for i in FECHADOS:
            cel(ws, r, 2 + i, f'=IF({L(1 + i)}{r0 + 1}=0,"",{L(2 + i)}{r0 + 1}/{L(1 + i)}{r0 + 1}-1)' if i >= 1 else None, VAR)
        r += 1
        cel(ws, r, 1, 'Δ share vs mesmo mês do ano anterior (pp)')
        dsh = lambda x, y: f'=IF(OR({x}{rs}="",{y}{rs}=""),"",({x}{rs}-{y}{rs})*100)'
        for i in FECHADOS:
            cel(ws, r, 2 + i, dsh(L(2 + i), L(2 + i - 12)) if i >= 12 else None, PP)
        cel(ws, r, NC + 2, dsh(L(NC + 2), L(NC + 3)), PP, bold=True)
        r += 1
        ws.conditional_formatting.add(f'B{rs}:{L(NC + 1)}{rs}',
                                      ColorScaleRule(start_type='min', start_color='FFFFFF', end_type='max', end_color=DEST))
        if t == 'REDE':
            ch = LineChart()
            ch.title = f'Rede · share {EMP} mês a mês'
            ch.height, ch.width = 6, 24
            ch.y_axis.number_format = '0%'
            ch.add_data(Reference(ws, min_col=1, max_col=NC + 1, min_row=rs, max_row=rs), from_rows=True, titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=2, max_col=NC + 1, min_row=r0 - 1))
            ws.add_chart(ch, f'B{r + 1}')
            r += 13
        r += 2
    larguras(ws, [34] + [11] * NC + [12, 12])
    ws.freeze_panes = 'B5'

    # =====================================================================================
    ws = wb.create_sheet('MARCAS E CATEGORIAS')
    titulo(ws, f'MARCAS DA EMPRESA E CATEGORIAS · rede e estados · {ACUM} {ANO} vs {ANO_ANT} · '
               f'{ml(fech)} vs {ml(M[AL])}', 'Share = empresa / mercado da categoria dentro da unidade.', 14)
    r = 4
    for t in UNIDADES:
        u = C[t]
        faixa(ws, r, u['name'] + ' · marcas da empresa', 14)
        r += 1
        cabecalho(ws, r, ['Marca', f'R$ {ACUM} {ANO}', f'R$ {ACUM} {ANO_ANT}', 'Var.', 'Peso na empresa',
                          f'R$ {ml(fech)}', f'R$ {ml(M[AL])}', f'Var. {MESES[mm - 1]}', f'Unid. {ACUM} {ANO}',
                          f'Unid. {ACUM} {ANO_ANT}', 'Var. unid.', f'Preço médio {ANO}', f'Preço médio {ANO_ANT}', ''])
        r += 1
        tr = S(u['emp']['v'], Y)
        for k, x in sorted(u['marca'].items(), key=lambda kv: -S(kv[1]['v'], Y)):
            v26, v25, q26, q25 = S(x['v'], Y), S(x['v'], YL), S(x['q'], Y), S(x['q'], YL)
            vals = [disp(k), v26, v25, var(v26, v25), share(v26, tr), x['v'][A], x['v'][AL], var(x['v'][A], x['v'][AL]),
                    q26, q25, var(q26, q25), v26 / q26 if q26 else None, v25 / q25 if q25 else None]
            fm = [None, BRL, BRL, VAR, PCT, BRL, BRL, VAR, INT, INT, VAR, '#,##0.00', '#,##0.00']
            for j, v in enumerate(vals):
                cel(ws, r, 1 + j, v, fm[j], bold=True, fill=EMPF)
            r += 1
        r += 1
        cabecalho(ws, r, ['Categoria', f'Mercado {ANO}', f'Mercado {ANO_ANT}', 'Var. mercado', f'{EMP} {ANO}',
                          f'{EMP} {ANO_ANT}', f'Var. {EMP}', f'Share {ANO}', f'Share {ANO_ANT}', 'Δ pp',
                          f'Share {ml(fech)}', f'Share {ml(M[AL])}', f'Δ pp {MESES[mm - 1]}', 'Marcas da empresa'])
        r += 1
        for k, x in sorted(u['cat_mkt'].items(), key=lambda kv: -S(kv[1]['v'], Y)):
            if not vale_cat(k):
                continue
            rb = u['cat_emp'].get(k, zero)
            a26, a25, b26, b25 = S(x['v'], Y), S(x['v'], YL), S(rb['v'], Y), S(rb['v'], YL)
            sa, sal = share(rb['v'][A], x['v'][A]), share(rb['v'][AL], x['v'][AL])
            mb = ', '.join(disp(m) for m, v in sorted([(kk.split('|')[1], S(vv['v'], Y)) for kk, vv in u['cat_marca'].items()
                                                       if kk.split('|')[0] == k], key=lambda z: -z[1]) if v > 0)
            vals = [k, a26, a25, var(a26, a25), b26, b25, var(b26, b25) if b25 else None, share(b26, a26), share(b25, a25),
                    ((share(b26, a26) or 0) - (share(b25, a25) or 0)) * 100, sa, sal, ((sa or 0) - (sal or 0)) * 100, mb or '-']
            fm = [None, BRL, BRL, VAR, BRL, BRL, VAR, PCT, PCT, PP, PCT, PCT, PP, None]
            for j, v in enumerate(vals):
                cel(ws, r, 1 + j, v, fm[j])
            r += 1
        r += 2
    larguras(ws, [28, 13, 13, 9, 12, 12, 12, 9, 12, 12, 9, 11, 11, 26])
    ws.freeze_panes = 'B5'

    # =====================================================================================
    ws = wb.create_sheet('SKUs DA EMPRESA')
    titulo(ws, f'SKUs DA EMPRESA NA REDE · {ACUM} {ANO} vs {ANO_ANT} · lojas com venda · preço · ranking',
           f'Lojas = número de lojas com venda em {ml(fech)} (máximo {NL}) e em {ml(M[AL])}. Sinal: perdeu lojas, entrou, saiu.', 14)
    r = 4
    cabecalho(ws, r, ['#', 'Marca', 'Categoria', 'EAN', 'SKU', f'R$ {ACUM} {ANO}', f'R$ {ACUM} {ANO_ANT}', 'Var.',
                      f'R$ {ml(fech)}', f'Lojas {ml(fech)}', f'Lojas {ml(M[AL])}', f'Preço médio {ANO}',
                      'Peso na empresa', 'Sinal'])
    r += 1
    tr = S(c['emp']['v'], Y)
    for i, s in enumerate(sorted(c['skus'], key=lambda s: -S(s['v'], Y)), 1):
        v26, v25 = S(s['v'], Y), S(s['v'], YL)
        l26, l25 = s['lojas'][A], s['lojas'][AL]
        q26 = S(s['q'], Y)
        g = var(v26, v25)
        sig = ('novo em ' + str(ANO4) if v25 == 0 and v26 > 0 else
               ('saiu do mix' if v26 == 0 and v25 > 0 else
                ('perdeu lojas' if l26 < l25 - LIM['sku_perdeu_lojas'] else
                 ('cai forte' if g is not None and g < -LIM['sku_forte'] else
                  ('cresce forte' if g is not None and g > LIM['sku_forte'] else '')))))
        fill = XF if sig in ('saiu do mix', 'perdeu lojas', 'cai forte') else (VF if sig.startswith(('cresce', 'novo')) else None)
        vals = [i, disp(s['marca']), s['cat'], s['ean'], s['nome'], v26, v25, g, s['v'][A], l26, l25,
                v26 / q26 if q26 else None, share(v26, tr), sig]
        fm = [INT, None, None, None, None, BRL, BRL, VAR, BRL, INT, INT, '#,##0.00', PCT, None]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j], fill=fill if j == 13 else None)
        r += 1
    larguras(ws, [5, 11, 24, 15, 52, 13, 13, 9, 11, 9, 9, 10, 9, 14])
    ws.freeze_panes = 'F5'

    # =====================================================================================
    CATS = [k for k, _ in sorted(((k, sum(ltot(d, P2, k) for d in lojas.values()))
                                  for k in {s['cat'] for d in lojas.values() for s in d['skus'].values()}),
                                 key=lambda kv: -kv[1])]
    MARCAS_LOJA = [m for m, _ in sorted(((m, sum(s['v'][j] for d in lojas.values() for s in d['skus'].values()
                                                 if s['marca'] == m for j in P2))
                                         for m in {s['marca'] for d in lojas.values() for s in d['skus'].values()}),
                                        key=lambda kv: -kv[1])]
    ws = wb.create_sheet('LOJAS')
    ncol = 11 + 1 + len(MARCAS_LOJA) + 1
    titulo(ws, f'ABERTURA POR LOJA · {NL} lojas · sell out da empresa · {FULLTXT} ({NSEM} semanas fechadas)',
           f'P1 = {n} semanas de {P1TXT}; P2 = {n} semanas de {P2TXT}. Onde cresce e onde cai = P2 vs P1. Mix por '
           f'categoria e marca em P2. O painel de varejo não abre loja a loja no ano anterior. UF atribuída pelo '
           f'uf_map.json da rede ou, na falta dele, pelo nome e endereço da loja.', ncol)
    r = 4
    faixa(ws, r, f'Resumo por estado (soma das lojas, {ANO4}: P1 = {P1TXT}, P2 = {P2TXT})', ncol)
    r += 1
    cabecalho(ws, r, ['UF', 'Lojas', f'Total {NSEM} sem.', 'Peso na rede', 'P1', 'P2', 'P2 vs P1',
                      'Lojas que crescem >5%', 'Lojas que caem >5%'] + [''] * (ncol - 9))
    r += 1
    TOTAL = sum(ltot(d, FULL) for d in lojas.values())
    ufs_lojas = sorted({uf_loja(k) for k in lojas}, key=lambda uf: -sum(ltot(d, FULL) for k, d in lojas.items() if uf_loja(k) == uf))
    for uf in ufs_lojas:
        ds = [(k, d) for k, d in lojas.items() if uf_loja(k) == uf]
        t_ = sum(ltot(d, FULL) for k, d in ds)
        p1 = sum(ltot(d, P1) for k, d in ds)
        p2 = sum(ltot(d, P2) for k, d in ds)
        up = sum(1 for k, d in ds if ltot(d, P1) and ltot(d, P2) / ltot(d, P1) - 1 > LIM['loja_movimento'])
        dn = sum(1 for k, d in ds if ltot(d, P1) and ltot(d, P2) / ltot(d, P1) - 1 < -LIM['loja_movimento'])
        vals = [uf, len(ds), t_, share(t_, TOTAL), p1, p2, var(p2, p1), up, dn]
        fm = [None, INT, BRL, PCT, BRL, BRL, VAR, INT, INT]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j], bold=True)
        r += 1
    r += 1
    cabecalho(ws, r, ['#', 'Loja', 'UF', 'Endereço', f'Total {NSEM} sem.', 'R$/semana', 'Peso na rede',
                      f'P1 ({n} sem.)', f'P2 ({n} sem.)', 'P2 vs P1', 'SKUs com venda P2', f'{CATS[0]} %']
              + [f'{disp(m)} %' for m in MARCAS_LOJA] + ['Movimento'])
    r += 1
    ws.freeze_panes = f'C{r}'
    r0 = r
    for i, (k, d) in enumerate(rk, 1):
        tot, p1, p2 = ltot(d, FULL), ltot(d, P1), ltot(d, P2)
        nsku = len([s for s in d['skus'].values() if sum(s['v'][j] for j in P2) > 0])
        mix_m = [(sum(s['v'][j] for s in d['skus'].values() if s['marca'] == m for j in P2) / p2) if p2 else 0
                 for m in MARCAS_LOJA]
        g = var(p2, p1)
        mv = 'cresce' if g is not None and g > LIM['loja_movimento'] else ('cai' if g is not None and g < -LIM['loja_movimento'] else 'estável')
        vals = [i, k, uf_loja(k), d['end'], tot, tot / NSEM, share(tot, TOTAL), p1, p2, g, nsku,
                share(ltot(d, P2, CATS[0]), p2)] + mix_m + [mv]
        fm = [INT, None, None, None, BRL, BRL, PCT, BRL, BRL, VAR, INT, PCT] + [PCT] * len(mix_m) + [None]
        for j, v in enumerate(vals):
            cel(ws, r, 1 + j, v, fm[j],
                fill=(XF if mv == 'cai' else (VF if mv == 'cresce' else None)) if j in (9, len(vals) - 1) else None)
        r += 1
    cel(ws, r, 2, 'Total rede', bold=True)
    for col in (5, 8, 9):
        cel(ws, r, col, f'=SUM({L(col)}{r0}:{L(col)}{r - 1})', BRL, bold=True)
    cel(ws, r, 10, f'=IF(H{r}=0,"",I{r}/H{r}-1)', VAR, bold=True)
    larguras(ws, [4, 36, 5, 34, 13, 11, 9, 12, 12, 9, 9, 12] + [9] * len(MARCAS_LOJA) + [10])

    # =====================================================================================
    ws = wb.create_sheet('LOJAS x SEMANA')
    titulo(ws, f'R$ DA EMPRESA POR LOJA E SEMANA · {W[0]} a {W[-1]}' + ('' if a.incluir_ultima_semana else ' (última semana parcial)'),
           'Escala de cor por linha ajuda a ver a sazonalidade e as quedas por loja.', nW + 3)
    r = 4
    cabecalho(ws, r, ['Loja', 'UF'] + W + [f'Total {NSEM} sem.'])
    r += 1
    r0 = r
    for k, d in rk:
        cel(ws, r, 1, k, bold=True)
        cel(ws, r, 2, uf_loja(k))
        for i in range(nW):
            cel(ws, r, 3 + i, sum(s['v'][i] for s in d['skus'].values()), BRL)
        cel(ws, r, nW + 3, f'=SUM({L(3 + FULL[0])}{r}:{L(3 + FULL[-1])}{r})', BRL, bold=True)
        r += 1
    cel(ws, r, 1, 'Total rede', bold=True)
    for i in range(nW + 1):
        cel(ws, r, 3 + i, f'=SUM({L(3 + i)}{r0}:{L(3 + i)}{r - 1})', BRL, bold=True)
    ws.conditional_formatting.add(f'C{r0}:{L(nW + 2)}{r - 1}',
                                  ColorScaleRule(start_type='min', start_color='FFFFFF', end_type='max', end_color=DEST))
    larguras(ws, [36, 5] + [10] * nW + [13])
    ws.freeze_panes = 'C5'

    # =====================================================================================
    ws = wb.create_sheet('LOJAS x CATEGORIA')
    titulo(ws, f'R$ DA EMPRESA POR LOJA E CATEGORIA · P2 ({P2TXT}) e variação vs P1',
           'Vermelho = categoria em queda acima de 10% na loja; cinza = categoria sem venda na loja (gap de sortimento).',
           3 + 2 * len(CATS))
    r = 4
    cabecalho(ws, r, ['Loja', 'UF', 'Total P2'] + [x for k in CATS for x in (k, 'var vs P1')])
    r += 1
    for k, d in rk:
        cel(ws, r, 1, k, bold=True)
        cel(ws, r, 2, uf_loja(k))
        cel(ws, r, 3, ltot(d, P2), BRL, bold=True)
        for j, cat in enumerate(CATS):
            x2, x1 = ltot(d, P2, cat), ltot(d, P1, cat)
            g = var(x2, x1)
            cel(ws, r, 4 + 2 * j, x2, BRL, fill=GF if x2 == 0 else None)
            cel(ws, r, 5 + 2 * j, g, VAR, fill=XF if g is not None and g < -LIM['categoria_loja'] else
                (VF if g is not None and g > LIM['categoria_loja'] else None))
        r += 1
    larguras(ws, [36, 5, 11] + [11, 8] * len(CATS))
    ws.freeze_panes = 'D5'

    # =====================================================================================
    ws = wb.create_sheet('LOJAS x SKU')
    titulo(ws, f'SORTIMENTO · top 30 SKUs da empresa na rede x {NL} lojas · R$ em P2 ({n} semanas)',
           'Célula cinza = SKU sem venda na loja no período: gap de sortimento a tratar. Última coluna = lojas sem venda por SKU.',
           4 + NL)
    sk = {}
    for k, d in lojas.items():
        for ean, s in d['skus'].items():
            x = sk.setdefault(ean, dict(nome=s['nome'], marca=s['marca'], tot=0))
            x['tot'] += sum(s['v'][i] for i in P2)
    top = sorted(sk.items(), key=lambda kv: -kv[1]['tot'])[:30]
    r = 4
    cabecalho(ws, r, ['SKU', 'Marca', 'Rede P2'] + [SNS(k) for k, _ in rk] + ['Lojas sem venda'])
    r += 1
    for ean, x in top:
        cel(ws, r, 1, x['nome'][:55], bold=True)
        cel(ws, r, 2, disp(x['marca']))
        cel(ws, r, 3, x['tot'], BRL)
        gaps = 0
        for j, (k, d) in enumerate(rk):
            s = d['skus'].get(ean)
            v = sum(s['v'][i] for i in P2) if s else 0
            cel(ws, r, 4 + j, v, BRL, fill=GF if v == 0 else None)
            gaps += 1 if v == 0 else 0
        cel(ws, r, 4 + len(rk), gaps, INT, bold=True, fill=XF if gaps >= LIM['gap_lojas_alerta'] else None)
        r += 1
    larguras(ws, [50, 11, 11] + [9] * len(rk) + [9])
    ws.freeze_panes = 'D5'

    # =====================================================================================
    cc = c.get('conc', {})
    if HASCONC:
        ws = wb.create_sheet('CONCORRÊNCIA')
        titulo(ws, f'CONCORRENTES DENTRO DA REDE · top marcas por categoria · {ACUM} {ANO} vs {ANO_ANT}',
               'Só existe quando a rede libera a concorrência em detalhe no painel de varejo.', 10)
        r = 4
        cabecalho(ws, r, ['Categoria', 'Marca', 'Fabricante', f'R$ {ACUM} {ANO}', f'R$ {ACUM} {ANO_ANT}', 'Var.',
                          f'Share na categoria {ANO}', f'Share {ANO_ANT}', 'Δ pp', 'Posição'])
        r += 1
        for cat, x in sorted(c['cat_mkt'].items(), key=lambda kv: -S(kv[1]['v'], Y)):
            if not vale_cat(cat):
                continue
            T26c, T25c = S(x['v'], Y), S(x['v'], YL)
            linhas_c = ([(disp(k.split('|')[1]), S(v['v'], Y), S(v['v'], YL), EMP) for k, v in c['cat_marca'].items()
                         if k.split('|')[0] == cat]
                        + [(k.split('|')[1], S(v['v'], Y), S(v['v'], YL), '') for k, v in cc.items() if k.split('|')[0] == cat])
            for i, (mk_, v26, v25, fab) in enumerate(sorted(linhas_c, key=lambda z: -z[1])[:8], 1):
                da_emp = fab == EMP
                vals = [cat, mk_, fab, v26, v25, var(v26, v25), share(v26, T26c), share(v25, T25c),
                        ((share(v26, T26c) or 0) - (share(v25, T25c) or 0)) * 100, i]
                fm = [None, None, None, BRL, BRL, VAR, PCT, PCT, PP, INT]
                for j, v in enumerate(vals):
                    cel(ws, r, 1 + j, v, fm[j], bold=da_emp, fill=EMPF if da_emp else None)
                r += 1
            r += 1
        larguras(ws, [26, 22, 12, 13, 13, 9, 12, 10, 9, 8])
        ws.freeze_panes = 'C5'

    # =====================================================================================
    ws = wb.create_sheet('PLANO DE AÇÃO')
    titulo(ws, f'PLANO DE AÇÃO · {REDE} · dono: {EXEC} · gerado a partir do diagnóstico do painel de varejo',
           'Cada linha nasce de um dado da planilha (coluna Evidência). Prazos sugeridos; ajustar com o cliente.', 7)
    r = 4
    cabecalho(ws, r, ['#', 'Frente', 'Ação', 'Evidência (painel de varejo)', 'Lojas / SKUs alvo', 'Dono', 'Prazo'])
    r += 1
    acoes = []

    def marcas_cat(k):
        return ', '.join(sorted({disp(kk.split('|')[1]) for kk in c['cat_marca']
                                 if kk.split('|')[0] == k and S(c['cat_marca'][kk]['v'], Y) > 0})) or 'da empresa'

    for x in PERDAS[:4]:
        k = x['cat']
        padrao = {'frente': k, 'acao': f'Recuperar {k}: revisar sortimento, preço de prateleira e execução em loja'}
        ac = ACOES.get(sem_acento(k), padrao)
        acoes.append((ac.get('frente', k), ac.get('acao', padrao['acao']),
                      f"{k}: mercado da rede {sp(var(x['mkt26'], x['mkt25']))} e empresa {sp(var(x['rb26'], x['rb25']))}; "
                      f"share de {x['s25'] * 100:.0f}% para {x['s26'] * 100:.0f}%",
                      f'SKUs {marcas_cat(k)} com sinal "cai forte" ou "perdeu lojas" (aba SKUs DA EMPRESA); '
                      f'lojas com gap na aba LOJAS x SKU', EXEC, mes_mais(2)))
    # presença: marcas com SKUs relevantes em menos lojas que a rede
    for mk in sorted({s['marca'] for s in c['skus']}, key=lambda m: -sum(S(s['v'], Y) for s in c['skus'] if s['marca'] == m)):
        sks = [s for s in c['skus'] if s['marca'] == mk and S(s['v'], Y) > 0]
        if not sks:
            continue
        tot = sum(S(s['v'], Y) for s in sks)
        rel = [s for s in sks if S(s['v'], Y) >= LIM['sku_relevante'] * tot]
        lj = [s['lojas'][A] for s in rel if s['lojas'][A] > 0]
        if lj and min(lj) <= NL - 2:
            acoes.append((f'Presença {disp(mk)}',
                          f'{disp(mk)}: garantir presença dos SKUs relevantes nas {NL} lojas (hoje {min(lj):.0f} a {max(lj):.0f} lojas por SKU)',
                          f"{len([s for s in rel if 0 < s['lojas'][A] < NL])} de {len(rel)} SKUs relevantes com venda em menos de "
                          f'{NL} lojas em {ml(fech)}',
                          f'SKUs {disp(mk)} com lojas < {NL} (aba SKUs DA EMPRESA, coluna Lojas {ml(fech)})', EXEC, mes_mais(2)))
    for k, p2, p1 in queda[-8:]:
        d = lojas[k]
        pior = sorted([(cat, ltot(d, P2, cat) - ltot(d, P1, cat)) for cat in CATS], key=lambda z: z[1])[:2]
        acoes.append(('Loja em queda', f'Visitar {SN(k)} ({uf_loja(k)}): checar ruptura, espaço e preço nas categorias que mais caíram',
                      f"P2 vs P1 {(p2 / p1 - 1) * 100:+.1f}% (R$ {BR(p1)} para R$ {BR(p2)}); maiores quedas: "
                      + '; '.join(f'{cat} (R$ {BR(v)})' for cat, v in pior), k, EXEC, mes_mais(1)))
    gapskus = [(x['nome'][:45], sum(1 for k, d in rk if (ean not in d['skus']) or sum(d['skus'][ean]['v'][i] for i in P2) == 0))
               for ean, x in top]
    gapskus = [g for g in gapskus if g[1] >= max(2, round(NL * LIM['gap_lojas_plano']))]
    if gapskus:
        acoes.append(('Sortimento', 'Fechar gaps de sortimento dos SKUs líderes nas lojas sem venda (pedido mínimo, cadastro, espaço)',
                      '; '.join(f'{nm}: {g} lojas' for nm, g in gapskus[:8]), 'Aba LOJAS x SKU (células cinza)', EXEC, mes_mais(2)))
    if GANHOS:
        acoes.append(('Defender o motor', f"{', '.join(x['cat'] for x in GANHOS[:3])}: manter execução (ponto extra, preço de "
                                          'prateleira, sortimento) onde o share já sobe',
                      '; '.join(f"{x['cat']} share {x['s26'] * 100:.0f}% ({x['dpp']:+.1f} pp)" for x in GANHOS[:3]),
                      'Todas as lojas', EXEC, 'contínuo'))
    if SEMEMP:
        x = SEMEMP[0]
        acoes.append((x['cat'], f"Levar ao cliente a decisão sobre a categoria {x['cat']} (sem presença relevante da empresa na rede)",
                      f"{x['cat']}: R$ {mi(x['mkt26'])} Mi {ACUM}/{ANO} ({sp(var(x['mkt26'], x['mkt25']))}), "
                      f"{x['mkt26'] / m26 * 100:.0f}% do mercado da rede, share da empresa {x['s26'] * 100:.1f}%",
                      'Pauta do plano de negócios com o cliente', GERENTE, mes_mais(3)))
    for i, ac in enumerate(acoes, 1):
        for j, v in enumerate([i] + list(ac)):
            cel(ws, r, 1 + j, ptbr(v) if isinstance(v, str) else v, wrap=True)
        ws.row_dimensions[r].height = 48
        r += 1
    larguras(ws, [4, 22, 52, 60, 44, 14, 10])

    # =====================================================================================
    ws = wb.create_sheet('FONTE')
    titulo(ws, 'FONTE E MÉTODO', None, 6)
    estados_exp = ', '.join(exportados) or 'nenhum'
    textos = [
        f'Portal do painel de varejo · relatório de KPIs por SKU (total da empresa): rede = {REDE}; filtro de '
        f'compartilhamento de dados do cliente = (Tudo)'
        + (' [a rede libera a concorrência por marca]' if HASCONC else
           ' [a rede libera só a empresa em detalhe e a concorrência agregada como "OUTRA MARCA / OUTROS PRODUTOS" por categoria]')
        + f'; unidade, segmento e categoria = (Tudo); top SKU cobrindo todo o universo; exports por estado do PDV '
          f'({estados_exp}). Meses {ml(M[0])} a {ml(M[-1])}' + (f'; {ml(PARCIAL)} parcial excluído.' if PARCIAL else '.'),
        f'Portal do painel de varejo · análise por rede, abertura por loja: rede = {REDE}, filtro de compartilhamento = '
        f'(Tudo), unidade = (Tudo), semanas de {W[0]} a {W[-1]}, {NL} lojas x SKU x semana. O portal não abre lojas antes '
        f'de {W[0]}.',
        'Mercado = todas as marcas nas categorias onde a empresa atua na rede'
        + (' (marcas concorrentes abertas na aba CONCORRÊNCIA)' if HASCONC else ' (o cliente não libera marcas concorrentes, só o agregado)')
        + f". A empresa = {' + '.join(MARCAS)}."
        + (f" {', '.join(FORA)} não é da empresa e conta como mercado." if FORA else '')
        + ' Nenhuma base própria do cliente nem Sell In foi usada neste arquivo.',
        f'Comparabilidade: lojas com venda por mês estão na aba MÊS A MÊS. {COMPTXT} Loja a loja: P1 ({P1TXT}) x P2 '
        f'({P2TXT}) de {ANO4}.',
        (f'Estados: {exportados[0]} exportado diretamente; demais estados = rede menos {exportados[0]} (aritmética exata do '
         f'próprio painel de varejo, sem estimativa)' if 'UF_DEMAIS' in C else
         'Estados exportados diretamente do painel de varejo (filtro de estado do PDV)')
        + f'; a abertura de {ANO4} por estado, loja a loja, está na aba LOJAS. Extraído em '
          f'{datetime.date.today().strftime("%d/%m/%Y")}. Gerado pela skill diagnostico-de-rede (montar_diagnostico.py).']
    for i, t in enumerate(textos, 3):
        paragrafo(ws, i, t, 6, 48)
    larguras(ws, [40, 20, 20, 20, 20, 20])

    saida = a.saida or os.path.join(a.pasta, 'DIAGNOSTICO_' + re.sub(r'\W+', '_', sem_acento(REDE)).strip('_')
                                    + '_PAINEL_VAREJO_LOJAS.xlsx')
    wb.save(saida)
    print('ok', saida, [w.title for w in wb.worksheets])


if __name__ == '__main__':
    main()
