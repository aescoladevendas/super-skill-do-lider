# -*- coding: utf-8 -*-
"""
Gera um slide de "conteúdo" (etiqueta + título + 4 KPI cards + tabela + caixa LEITURA)
no padrão visual rosa/slate da skill book-executivo.

Entradas: os textos e números do slide (parâmetros de build_slide, slide_xml ou
escrever_no_slide). Cores e fonte vêm de saidas.pptx_executivo do arquivo de configuração
(com saidas.identidade_visual.destaque e fonte_office como reserva) e têm padrão no código.
Saídas: o XML completo do slide (arquivo slideN.xml) ou o slide escrito direto num .pptx.

Dois caminhos de uso:

1) Fluxo descompactar, editar e reempacotar (unpack, clean e pack com a skill pptx do Claude):

    import sys
    sys.path.insert(0, "<pasta da skill>/scripts")
    from build_content_slide import build_slide

    build_slide(
        eyebrow="SELL IN",
        title="Desempenho por executivo (P07/26)",
        page="06/11",
        footer="Empresa Exemplo · Book T3 2026 · Regional Exemplo",
        cards_spec=[
            {"label": "REAL NR (FATURADO)", "value": "R$ 2,16 Mi", "delta": "7,5% da meta", "deltacolor": "negativo"},
            {"label": "GAP PÓS-CARTEIRA", "value": "R$ 24,01 Mi", "delta": "Necessário até set/26", "deltacolor": "neutro"},
            {"label": "MAIOR CONCENTRAÇÃO", "value": "Luísa Andrade: 57%", "delta": "Risco de concentração", "deltacolor": "negativo", "valsz": 1500},
            {"label": "PIOR DESEMPENHO", "value": "Executivo C: -0,3%", "delta": "NR negativo no mês", "deltacolor": "negativo", "valsz": 1500},
        ],
        headers=["Executivo", "Canal", "Meta NR", "Real NR", "% Atg", "Carteira NR", "Fech. Proj.", "% Fech Proj"],
        data_rows=[
            ("Executivo B", "Direto", "R$ 5.724.059", "R$ 492.357", "8,6%", "R$ 1.179.670", "R$ 1.672.027", "29,2%"),
            # ... uma linha por item; a ÚLTIMA linha é sempre o TOTAL
            ("TOTAL", "-", "R$ 28.722.385", "R$ 2.156.849", "7,5%", "R$ 2.552.560", "R$ 4.709.409", "16,4%"),
        ],
        leitura="Frase de diagnóstico + insight: no MÁXIMO 2 linhas, senão colide com o rodapé.",
        out_path="descompactado/ppt/slides/slideN.xml",
    )

2) Direto no .pptx com python-pptx (sem descompactar):

    from pptx import Presentation
    from build_content_slide import escrever_no_slide
    prs = Presentation("template-base.pptx")
    escrever_no_slide(prs.slides[3], eyebrow=..., title=..., ...)   # mesmos campos, sem out_path
    prs.save("book.pptx")

"deltacolor" aceita um hexadecimal ("C0392B") ou o nome da cor: "negativo", "positivo",
"neutro" (slate). Para trocar uma cor só nesta execução: PALETA["acento"] = "0071E3".

Teste rápido (escreve um slide de demonstração com dados fictícios):
    python3 build_content_slide.py demo_slide.xml
    python3 build_content_slide.py --pptx ../assets/template-base.pptx book_teste.pptx 4

Lição aprendida (não pule esta parte): cada linha de `data_rows` (inclusive a linha TOTAL)
precisa ter EXATAMENTE o mesmo número de colunas que `headers`. Se uma linha tiver menos
células que `<a:gridCol>` definidos na tabela, PowerPoint e LibreOffice "improvisam" a célula
que falta com um estilo de reserva (apareceu como um retângulo azul sólido cobrindo cerca de 2
linhas de altura na primeira tentativa) em vez de dar erro; o defeito só aparece no render, não
na validação de XML. Por isso `build_table` confere o tamanho de cada linha e falha alto
(AssertionError) se algo não bater, e a grade da tabela é montada com o número de colunas de
`headers`.
"""
import os
import sys
from xml.sax.saxutils import escape

AQUI = os.path.dirname(os.path.abspath(__file__))


def _paleta():
    """Cores sem '#' e fonte, lidas da configuração (se houver) com padrão no código."""
    pal = {"acento": "FF007F", "slate": "415A6C", "slate_medio": "64748B", "slate_claro": "94A3B8",
           "fundo_conteudo": "F8FAFC", "zebra": "F1F5F9", "borda": "E2E8F0",
           "positivo": "50B428", "negativo": "C0392B", "fonte": "Arial"}
    try:
        sys.path.insert(0, AQUI)
        from configuracao import carregar, obter
        cfg = carregar()
    except (Exception, SystemExit):
        return pal
    destaque = obter(cfg, "saidas.identidade_visual.destaque")
    if destaque:
        pal["acento"] = str(destaque).lstrip("#")
    fonte = obter(cfg, "saidas.identidade_visual.fonte_office")
    if fonte:
        pal["fonte"] = fonte
    for k, v in (obter(cfg, "saidas.pptx_executivo", {}) or {}).items():
        if k in pal and v:
            pal[k] = str(v).lstrip("#")
    return pal


PALETA = _paleta()


def _cor(c):
    """aceita hexadecimal ou nome semântico (positivo, negativo, neutro, acento...)."""
    nomes = {"neutro": "slate"}
    c = str(c).lstrip("#")
    return PALETA.get(nomes.get(c, c), c)


def _t(texto):
    return escape(str(texto))


def cell(text, color=None, algn="r", bold=False, bg=None, sz=950, mar=27432):
    color = _cor(color or "slate")
    bg = _cor(bg or "FFFFFF")
    b = ' b="1"' if bold else ''
    return f'''                <a:tc>
                  <a:txBody>
                    <a:bodyPr/>
                    <a:lstStyle/>
                    <a:p>
                      <a:pPr algn="{algn}"/>
                      <a:r>
                        <a:rPr sz="{sz}"{b}>
                          <a:solidFill>
                            <a:srgbClr val="{color}"/>
                          </a:solidFill>
                          <a:latin typeface="{PALETA['fonte']}"/>
                        </a:rPr>
                        <a:t>{_t(text)}</a:t>
                      </a:r>
                    </a:p>
                  </a:txBody>
                  <a:tcPr marT="{mar}" marB="{mar}" anchor="ctr">
                    <a:solidFill>
                      <a:srgbClr val="{bg}"/>
                    </a:solidFill>
                  </a:tcPr>
                </a:tc>'''


def row(vals, row_id, bold_row=False, band=False, header=False, row_h=265176, sz=950, mar=27432):
    bg = "slate" if header else ("zebra" if band else "FFFFFF")
    color = "FFFFFF" if header else "slate"
    cells = [cell(t, color=color, algn=a, bold=(header or bold_row), bg=bg, sz=sz, mar=mar) for t, a in vals]
    return f'''              <a:tr h="{row_h}">
{chr(10).join(cells)}
                <a:extLst>
                  <a:ext uri="{{0D108BD9-81ED-4DB2-BD59-A6C34878D82A}}">
                    <a16:rowId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" val="{row_id}"/>
                  </a:ext>
                </a:extLst>
              </a:tr>'''


# Altura disponível real entre o topo da tabela (y=2971800) e o topo da caixa LEITURA
# (y=5486400) no template: nunca deixar a tabela ultrapassar isso, senão invade a LEITURA.
TABLE_MAX_HEIGHT = 2514600
TABLE_WIDTH = 11183112
# grade padrão de 8 colunas: 1 larga para nome/ID, 6 médias de dado, 1 larga para leitura textual
GRID_8 = [2377440, 1097280, 1097280, 1097280, 1097280, 1097280, 1097280, 2221992]


def table_sizing(n_rows):
    """Escolhe row_h/sz/mar para caber n_rows (cabeçalho + dados + total) sem estourar
    TABLE_MAX_HEIGHT. Até 9 linhas usa o tamanho padrão; tabelas maiores (ex.: lista completa
    de segmentos, sem cortar nos 6 primeiros) compactam fonte, altura e margem aos poucos.
    Testado visualmente até cerca de 16 linhas."""
    if n_rows <= 9:
        return 265176, 950, 27432
    elif n_rows <= 12:
        return 195000, 900, 18288
    elif n_rows <= 16:
        return 155000, 800, 9144
    else:
        return max(120000, TABLE_MAX_HEIGHT // n_rows), 700, 4572


def grid_widths(n_cols):
    """larguras das colunas: a grade padrão de 8; com outro número de colunas, primeira e
    última largas e o meio dividido por igual, sempre somando TABLE_WIDTH."""
    if n_cols == 8:
        return list(GRID_8)
    if n_cols == 1:
        return [TABLE_WIDTH]
    if n_cols == 2:
        return [GRID_8[0], TABLE_WIDTH - GRID_8[0]]
    meio = (TABLE_WIDTH - GRID_8[0] - GRID_8[-1]) // (n_cols - 2)
    w = [GRID_8[0]] + [meio] * (n_cols - 2)
    return w + [TABLE_WIDTH - sum(w)]


def build_grid(n_cols):
    cols = []
    for i, w in enumerate(grid_widths(n_cols)):
        cols.append(f'                <a:gridCol w="{w}"><a:extLst><a:ext uri="{{9D8B030D-6E8A-4147-A177-3AD203B41FA5}}">'
                    f'<a16:colId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" val="{20000 + i}"/></a:ext></a:extLst></a:gridCol>')
    return "\n".join(cols)


def build_table(headers, data_rows, row_h=265176, sz=950, mar=27432):
    n_cols = len(headers)
    for i, r in enumerate(data_rows):
        assert len(r) == n_cols, (
            f"data_rows[{i}] tem {len(r)} colunas, mas headers tem {n_cols}. "
            "Toda linha (inclusive TOTAL) precisa ter o mesmo número de colunas "
            "que headers, senão a tabela renderiza uma célula 'fantasma' com "
            "cor de reserva em vez de dar erro de XML."
        )
    hdr = row([(h, "l" if i == 0 else "r") for i, h in enumerate(headers)], "10000", header=True, row_h=row_h, sz=sz, mar=mar)
    body_rows = []
    n = len(data_rows)
    for i, r in enumerate(data_rows):
        is_total = (i == n - 1)
        vals = [(str(v), "l" if j == 0 else "r") for j, v in enumerate(r)]
        body_rows.append(row(vals, str(10001 + i), bold_row=is_total, band=(i % 2 == 1) or is_total, row_h=row_h, sz=sz, mar=mar))
    return hdr + "\n" + "\n".join(body_rows)


KPI_TEMPLATE = '''      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="{id1}" name="Rectangle {rid1}"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{x}" y="1600200"/>
            <a:ext cx="2542032" cy="1234440"/>
          </a:xfrm>
          <a:prstGeom prst="rect">
            <a:avLst/>
          </a:prstGeom>
          <a:solidFill>
            <a:srgbClr val="FFFFFF"/>
          </a:solidFill>
          <a:ln w="9525">
            <a:solidFill>
              <a:srgbClr val="{borda}"/>
            </a:solidFill>
          </a:ln>
          <a:effectLst/>
        </p:spPr>
        <p:style>
          <a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef>
          <a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>
          <a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>
          <a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>
        </p:style>
        <p:txBody>
          <a:bodyPr rtlCol="0" anchor="ctr"/>
          <a:lstStyle/>
          <a:p><a:pPr algn="ctr"/><a:endParaRPr/></a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="{id2}" name="Rectangle {rid2}"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{x}" y="1600200"/>
            <a:ext cx="64008" cy="1234440"/>
          </a:xfrm>
          <a:prstGeom prst="rect">
            <a:avLst/>
          </a:prstGeom>
          <a:solidFill>
            <a:srgbClr val="{acento}"/>
          </a:solidFill>
          <a:ln><a:noFill/></a:ln>
          <a:effectLst/>
        </p:spPr>
        <p:style>
          <a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef>
          <a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>
          <a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>
          <a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>
        </p:style>
        <p:txBody>
          <a:bodyPr rtlCol="0" anchor="ctr"/>
          <a:lstStyle/>
          <a:p><a:pPr algn="ctr"/><a:endParaRPr/></a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="{id3}" name="TextBox {rid3}"/>
          <p:cNvSpPr txBox="1"/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{xl}" y="1764792"/>
            <a:ext cx="2130552" cy="320040"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="1050" b="1" i="0" spc="200">
                <a:solidFill><a:srgbClr val="{slate_medio}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{label}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="{id4}" name="TextBox {rid4}"/>
          <p:cNvSpPr txBox="1"/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{xl}" y="2075688"/>
            <a:ext cx="2130552" cy="{valh}"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"><a:lnSpc><a:spcPct val="105000"/></a:lnSpc></a:pPr>
            <a:r>
              <a:rPr sz="{valsz}" b="1" i="0">
                <a:solidFill><a:srgbClr val="{slate}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{value}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="{id5}" name="TextBox {rid5}"/>
          <p:cNvSpPr txBox="1"/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{xl}" y="{deltay}"/>
            <a:ext cx="2130552" cy="292608"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="1100" b="1" i="0">
                <a:solidFill><a:srgbClr val="{deltacolor}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{delta}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>'''


def build_kpi_card(idx, x, label, value, delta, deltacolor, valsz=2000, valh=566928, deltay=2468880):
    ids = list(range(6 + idx * 5, 11 + idx * 5))
    return KPI_TEMPLATE.format(
        id1=ids[0], rid1=ids[0] - 1, id2=ids[1], rid2=ids[1] - 1, id3=ids[2], rid3=ids[2] - 1,
        id4=ids[3], rid4=ids[3] - 1, id5=ids[4], rid5=ids[4] - 1,
        x=x, xl=x + 228600, label=_t(label), value=_t(value), valsz=valsz, valh=valh, deltay=deltay,
        delta=_t(delta), deltacolor=_cor(deltacolor), **PALETA
    )


HEAD_TEMPLATE = '''<?xml version="1.0" encoding="utf-8"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:bg>
      <p:bgPr>
        <a:solidFill><a:srgbClr val="{fundo_conteudo}"/></a:solidFill>
        <a:effectLst/>
      </p:bgPr>
    </p:bg>
    <p:spTree>
      <p:nvGrpSpPr>
        <p:cNvPr id="1" name=""/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm>
          <a:off x="0" y="0"/><a:ext cx="0" cy="0"/>
          <a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/>
        </a:xfrm>
      </p:grpSpPr>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="2" name="TextBox 1"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="502920" y="411480"/><a:ext cx="10515600" cy="292608"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="1100" b="1" i="0" spc="400">
                <a:solidFill><a:srgbClr val="{acento}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{eyebrow}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="3" name="TextBox 2"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="502920" y="713232"/><a:ext cx="10515600" cy="777240"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="2800" b="1" i="0">
                <a:solidFill><a:srgbClr val="{slate}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{title}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="4" name="TextBox 3"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="502920" y="6547104"/><a:ext cx="7772400" cy="274320"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="900" b="0" i="0">
                <a:solidFill><a:srgbClr val="{slate_claro}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{footer}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="5" name="TextBox 4"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="10972800" y="6547104"/><a:ext cx="713232" cy="274320"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="r"/>
            <a:r>
              <a:rPr sz="900" b="0" i="0">
                <a:solidFill><a:srgbClr val="{slate_claro}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{page}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
{cards}
      <p:graphicFrame>
        <p:nvGraphicFramePr>
          <p:cNvPr id="26" name="Table 25"/>
          <p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr>
          <p:nvPr/>
        </p:nvGraphicFramePr>
        <p:xfrm>
          <a:off x="502920" y="2971800"/>
          <a:ext cx="{table_cx}" cy="{table_cy}"/>
        </p:xfrm>
        <a:graphic>
          <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">
            <a:tbl>
              <a:tblPr firstRow="1" bandRow="1">
                <a:tableStyleId>{{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}}</a:tableStyleId>
              </a:tblPr>
              <a:tblGrid>
{table_grid}
              </a:tblGrid>
{table_rows}
            </a:tbl>
          </a:graphicData>
        </a:graphic>
      </p:graphicFrame>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="27" name="Rectangle 26"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="502920" y="5486400"/><a:ext cx="11183112" cy="914400"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
          <a:ln w="9525"><a:solidFill><a:srgbClr val="{borda}"/></a:solidFill></a:ln>
          <a:effectLst/>
        </p:spPr>
        <p:style>
          <a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef>
          <a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>
          <a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>
          <a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>
        </p:style>
        <p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:endParaRPr/></a:p></p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="28" name="Rectangle 27"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="502920" y="5486400"/><a:ext cx="64008" cy="914400"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="{acento}"/></a:solidFill>
          <a:ln><a:noFill/></a:ln>
          <a:effectLst/>
        </p:spPr>
        <p:style>
          <a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef>
          <a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>
          <a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>
          <a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>
        </p:style>
        <p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:endParaRPr/></a:p></p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="29" name="TextBox 28"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="795528" y="5687568"/><a:ext cx="10634472" cy="274320"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"/>
            <a:r>
              <a:rPr sz="1100" b="1" i="0" spc="300">
                <a:solidFill><a:srgbClr val="{acento}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>LEITURA</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="30" name="TextBox 29"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
        <p:spPr>
          <a:xfrm><a:off x="795528" y="5998464"/><a:ext cx="10634472" cy="182880"/></a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:noFill/>
        </p:spPr>
        <p:txBody>
          <a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:spAutoFit/></a:bodyPr>
          <a:lstStyle/>
          <a:p>
            <a:pPr algn="l"><a:lnSpc><a:spcPct val="125000"/></a:lnSpc></a:pPr>
            <a:r>
              <a:rPr sz="1300" b="0" i="0">
                <a:solidFill><a:srgbClr val="{slate}"/></a:solidFill>
                <a:latin typeface="{fonte}"/>
              </a:rPr>
              <a:t>{leitura}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
  <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>
'''


def slide_xml(eyebrow, title, page, footer, cards_spec, headers, data_rows, leitura):
    """devolve o XML completo do slide de conteúdo (texto)."""
    xs = [502920, 3200400, 5897880, 8595360]
    cards_xml = "\n".join(
        build_kpi_card(i, xs[i], c["label"], c["value"], c["delta"], c.get("deltacolor", "neutro"),
                       c.get("valsz", 2000), c.get("valh", 566928), c.get("deltay", 2468880))
        for i, c in enumerate(cards_spec)
    )
    n_rows = len(data_rows) + 1  # +1 pela linha de cabeçalho
    row_h, sz, mar = table_sizing(n_rows)
    table_xml = build_table(headers, data_rows, row_h=row_h, sz=sz, mar=mar)
    table_cy = min(TABLE_MAX_HEIGHT, row_h * n_rows)
    return HEAD_TEMPLATE.format(eyebrow=_t(eyebrow), title=_t(title), page=_t(page), footer=_t(footer),
                                cards=cards_xml, table_grid=build_grid(len(headers)), table_rows=table_xml,
                                leitura=_t(leitura), table_cx=TABLE_WIDTH, table_cy=table_cy, **PALETA)


def build_slide(eyebrow, title, page, footer, cards_spec, headers, data_rows, leitura, out_path):
    """grava o slide de conteúdo em out_path (ex.: descompactado/ppt/slides/slide4.xml)."""
    xml = slide_xml(eyebrow, title, page, footer, cards_spec, headers, data_rows, leitura)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(xml)
    print("gravado", out_path, f"({len(data_rows)} linhas de dado + cabeçalho, {len(headers)} colunas)")


def escrever_no_slide(slide, **campos):
    """escreve o slide de conteúdo direto num slide aberto com python-pptx (sem descompactar)."""
    from pptx.oxml import parse_xml
    novo = parse_xml(slide_xml(**campos).encode("utf-8"))
    velho = slide._element
    for filho in list(velho):
        velho.remove(filho)
    for filho in list(novo):
        velho.append(filho)


def demo():
    """campos de um slide de demonstração, com os nomes da configuração em uso."""
    try:
        sys.path.insert(0, AQUI)
        from configuracao import carregar, nomes_executivos
        nomes = nomes_executivos(carregar())
    except (Exception, SystemExit):
        nomes = []
    nomes = (nomes + ["Executivo A", "Executivo B", "Executivo C"])[:3]
    return dict(
        eyebrow="SELL IN",
        title="Desempenho por executivo (P07/26)",
        page="04/11",
        footer="Empresa Exemplo · Book T3 2026 · Regional Exemplo",
        cards_spec=[
            {"label": "REAL NR (FATURADO)", "value": "R$ 2,16 Mi", "delta": "7,5% da meta", "deltacolor": "negativo"},
            {"label": "GAP PÓS-CARTEIRA", "value": "R$ 24,01 Mi", "delta": "Necessário até set/26", "deltacolor": "neutro"},
            {"label": "MAIOR CONCENTRAÇÃO", "value": f"{nomes[0]}: 57%", "delta": "Risco de concentração", "deltacolor": "negativo", "valsz": 1500},
            {"label": "PIOR DESEMPENHO", "value": f"{nomes[2]}: -0,1%", "delta": "NR negativo no mês", "deltacolor": "negativo", "valsz": 1500},
        ],
        headers=["Executivo", "Canal", "Meta NR", "Real NR", "% Atg", "Carteira NR", "Fech. Proj.", "% Fech Proj"],
        data_rows=[
            (nomes[0], "Indireto", "R$ 16.350.000", "R$ 1.229.404", "7,5%", "R$ 1.020.000", "R$ 2.249.404", "13,8%"),
            (nomes[1], "Direto", "R$ 5.724.059", "R$ 935.380", "16,3%", "R$ 1.179.670", "R$ 2.115.050", "37,0%"),
            (nomes[2], "Atacado", "R$ 6.648.326", "R$ -7.935", "-0,1%", "R$ 352.890", "R$ 344.955", "5,2%"),
            ("TOTAL", "-", "R$ 28.722.385", "R$ 2.156.849", "7,5%", "R$ 2.552.560", "R$ 4.709.409", "16,4%"),
        ],
        leitura="O mês fechou em 7,5% da meta e a carteira leva a projeção a 16,4%. P&L e devoluções explicam a "
                "queda do Atacado: recuperar o pedido parado é a ação da semana.",
    )


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) >= 3 and a[0] == "--pptx":
        from pptx import Presentation
        prs = Presentation(a[1])
        n = int(a[3]) if len(a) > 3 else 4
        escrever_no_slide(prs.slides[n - 1], **demo())
        prs.save(a[2])
        print("gravado", a[2], f"(slide {n} reescrito com a demonstração)")
    elif len(a) == 1:
        build_slide(out_path=a[0], **demo())
    else:
        print(__doc__)
